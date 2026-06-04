import io
from pathlib import Path
from redops.models.engagement import Engagement, Objective, TTP, LogEntry, OpsecResult


class DOCXReporter:
    def generate(
        self,
        engagement: Engagement,
        objectives: list[Objective],
        ttps: list[TTP],
        log_entries: list[LogEntry],
        opsec_results: list[OpsecResult] | None = None,
    ) -> bytes:
        from docx import Document
        from docx.shared import Pt, RGBColor
        from docx.enum.text import WD_ALIGN_PARAGRAPH

        doc = Document()

        # Cover
        title = doc.add_heading(engagement.name, 0)
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        doc.add_paragraph(f"Client: {engagement.client}")
        doc.add_paragraph(f"Period: {engagement.start_date} — {engagement.end_date}")
        doc.add_paragraph(f"Status: {engagement.status}")
        doc.add_paragraph(f"Operators: {', '.join(engagement.operators) or 'N/A'}")
        doc.add_page_break()

        # Executive Summary
        doc.add_heading("Executive Summary", 1)
        achieved = sum(1 for o in objectives if o.status == "achieved")
        doc.add_paragraph(f"Objectives achieved: {achieved}/{len(objectives)}")
        doc.add_paragraph(f"TTPs exercised: {len(ttps)}")
        doc.add_paragraph(f"Operator log entries: {len(log_entries)}")

        # Objectives
        doc.add_heading("Objectives", 1)
        for obj in objectives:
            p = doc.add_paragraph(style="List Bullet")
            p.add_run(f"[{obj.status.upper()}] {obj.title}").bold = True
            if obj.description:
                doc.add_paragraph(obj.description, style="Normal")

        # MITRE ATT&CK Coverage
        doc.add_heading("MITRE ATT&CK Coverage", 1)
        if ttps:
            table = doc.add_table(rows=1, cols=5)
            table.style = "Table Grid"
            hdr = table.rows[0].cells
            for i, h in enumerate(["ID", "Name", "Tactic", "Phase", "Status"]):
                hdr[i].text = h
            for ttp in sorted(ttps, key=lambda t: t.tactic):
                row = table.add_row().cells
                row[0].text = ttp.technique_id
                row[1].text = ttp.technique_name
                row[2].text = ttp.tactic
                row[3].text = ttp.phase
                row[4].text = ttp.status
        else:
            doc.add_paragraph("No TTPs recorded.")

        # Technical Narrative
        doc.add_heading("Technical Narrative", 1)
        if log_entries:
            table = doc.add_table(rows=1, cols=6)
            table.style = "Table Grid"
            hdr = table.rows[0].cells
            for i, h in enumerate(["Timestamp", "Operator", "Action", "Target", "Technique", "Result"]):
                hdr[i].text = h
            for entry in sorted(log_entries, key=lambda e: e.timestamp):
                row = table.add_row().cells
                row[0].text = entry.timestamp.strftime("%Y-%m-%d %H:%M")
                row[1].text = entry.operator
                row[2].text = entry.action
                row[3].text = entry.target
                row[4].text = entry.technique_id or ""
                row[5].text = entry.result
        else:
            doc.add_paragraph("No operator log entries recorded.")

        # Detection Gaps
        doc.add_heading("Detection Gaps", 1)
        gaps = [e for e in log_entries if e.result == "success"]
        doc.add_paragraph(f"{len(gaps)} action(s) succeeded without detection.")
        for e in gaps:
            doc.add_paragraph(f"• {e.action} on {e.target} ({e.technique_id or 'no technique'})", style="Normal")

        # Appendix
        doc.add_heading("Appendix — Scope & Rules of Engagement", 1)
        doc.add_heading("In Scope", 2)
        for s in engagement.scope:
            doc.add_paragraph(s, style="List Bullet")
        doc.add_heading("Out of Scope", 2)
        for s in engagement.out_of_scope:
            doc.add_paragraph(s, style="List Bullet")
        doc.add_heading("Rules of Engagement", 2)
        doc.add_paragraph(engagement.rules_of_engagement or "N/A")

        buf = io.BytesIO()
        doc.save(buf)
        return buf.getvalue()

    def write(self, engagement: Engagement, objectives: list[Objective],
              ttps: list[TTP], log_entries: list[LogEntry],
              path: str | Path, opsec_results: list[OpsecResult] | None = None) -> None:
        Path(path).write_bytes(
            self.generate(engagement, objectives, ttps, log_entries, opsec_results)
        )
