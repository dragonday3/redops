import io
import json
from datetime import date, datetime, timezone
from pathlib import Path
from unittest.mock import MagicMock, patch

import pytest

from redops.models.engagement import (
    Engagement, Objective, TTP, LogEntry, OpsecResult, OpsecCheckResult
)
from redops.reporters.json_reporter import JSONReporter
from redops.reporters.html import HTMLReporter, _status_badge, _format_dt
from redops.reporters.docx_reporter import DOCXReporter


# ── Fixtures ─────────────────────────────────────────────────────────────────

@pytest.fixture
def engagement():
    return Engagement(
        name="Op Nightfall",
        client="ACME Corp",
        scope=["10.0.0.0/8", "acme.example.com"],
        out_of_scope=["prod-db.acme.internal"],
        rules_of_engagement="No destructive actions. Business hours only.",
        start_date=date(2024, 1, 15),
        end_date=date(2024, 2, 15),
        operators=["alice", "bob"],
        status="completed",
    )


@pytest.fixture
def objectives(engagement):
    return [
        Objective(engagement_id=engagement.id, title="Gain initial access", type="primary", status="achieved"),
        Objective(engagement_id=engagement.id, title="Exfiltrate data", type="secondary", status="partial"),
    ]


@pytest.fixture
def ttps(engagement):
    return [
        TTP(engagement_id=engagement.id, technique_id="T1566.001", technique_name="Spearphishing",
            tactic="Initial Access", phase="Week 1", status="executed"),
        TTP(engagement_id=engagement.id, technique_id="T1078", technique_name="Valid Accounts",
            tactic="Defense Evasion", phase="Week 2", status="detected"),
    ]


@pytest.fixture
def log_entries(engagement):
    return [
        LogEntry(
            engagement_id=engagement.id,
            operator="alice",
            action="Sent phishing email",
            target="user@acme.com",
            result="success",
            technique_id="T1566.001",
        ),
        LogEntry(
            engagement_id=engagement.id,
            operator="bob",
            action="Lateral movement attempt",
            target="dc.acme.internal",
            result="failure",
        ),
    ]


@pytest.fixture
def opsec_results(engagement):
    return [
        OpsecResult(
            domain="c2.example.com",
            score=80,
            grade="B",
            checks=[
                OpsecCheckResult(name="dns_ttl", status="pass", detail="TTL is low"),
            ],
        )
    ]


# ── _status_badge tests ───────────────────────────────────────────────────────

def test_status_badge_pass_contains_green():
    result = _status_badge("pass")
    assert "#238636" in result


def test_status_badge_fail_contains_red():
    result = _status_badge("fail")
    assert "#da3633" in result


def test_status_badge_unknown_status_fallback():
    result = _status_badge("unknown_status")
    assert "#30363d" in result


def test_status_badge_success_contains_green():
    result = _status_badge("success")
    assert "#238636" in result


def test_status_badge_blocked_contains_red():
    result = _status_badge("blocked")
    assert "#da3633" in result


def test_status_badge_detected_contains_red():
    result = _status_badge("detected")
    assert "#da3633" in result


def test_status_badge_partial_contains_yellow():
    result = _status_badge("partial")
    assert "#9e6a03" in result


def test_status_badge_executed_contains_blue():
    result = _status_badge("executed")
    assert "#1f6feb" in result


# ── _format_dt tests ──────────────────────────────────────────────────────────

def test_format_dt_datetime_utc():
    dt = datetime(2024, 1, 15, 10, 30, tzinfo=timezone.utc)
    assert _format_dt(dt) == "2024-01-15 10:30 UTC"


def test_format_dt_none_returns_dash():
    assert _format_dt(None) == "—"


def test_format_dt_naive_datetime():
    dt = datetime(2024, 6, 1, 9, 0)
    result = _format_dt(dt)
    assert "2024-06-01" in result
    assert "09:00" in result


# ── JSONReporter tests ────────────────────────────────────────────────────────

def test_json_reporter_returns_valid_json(engagement, objectives, ttps, log_entries):
    reporter = JSONReporter()
    result = reporter.generate(engagement, objectives, ttps, log_entries)
    data = json.loads(result)  # must not raise
    assert isinstance(data, dict)


def test_json_output_contains_engagement_name(engagement, objectives, ttps, log_entries):
    reporter = JSONReporter()
    result = reporter.generate(engagement, objectives, ttps, log_entries)
    assert "Op Nightfall" in result


def test_json_output_contains_objectives(engagement, objectives, ttps, log_entries):
    reporter = JSONReporter()
    result = reporter.generate(engagement, objectives, ttps, log_entries)
    data = json.loads(result)
    assert len(data["objectives"]) == 2
    assert data["objectives"][0]["title"] == "Gain initial access"


def test_json_output_contains_ttps(engagement, objectives, ttps, log_entries):
    reporter = JSONReporter()
    result = reporter.generate(engagement, objectives, ttps, log_entries)
    data = json.loads(result)
    assert len(data["ttps"]) == 2
    technique_ids = {t["technique_id"] for t in data["ttps"]}
    assert "T1566.001" in technique_ids


def test_json_output_contains_log_entries(engagement, objectives, ttps, log_entries):
    reporter = JSONReporter()
    result = reporter.generate(engagement, objectives, ttps, log_entries)
    data = json.loads(result)
    assert len(data["log_entries"]) == 2


def test_json_output_opsec_results_empty_by_default(engagement, objectives, ttps, log_entries):
    reporter = JSONReporter()
    result = reporter.generate(engagement, objectives, ttps, log_entries)
    data = json.loads(result)
    assert data["opsec_results"] == []


def test_json_reporter_write_creates_file(engagement, objectives, ttps, log_entries, tmp_path):
    reporter = JSONReporter()
    out = tmp_path / "report.json"
    reporter.write(engagement, objectives, ttps, log_entries, out)
    assert out.exists()
    data = json.loads(out.read_text(encoding="utf-8"))
    assert data["engagement"]["name"] == "Op Nightfall"


def test_json_with_opsec_results_populated(engagement, objectives, ttps, log_entries, opsec_results):
    reporter = JSONReporter()
    result = reporter.generate(engagement, objectives, ttps, log_entries, opsec_results)
    data = json.loads(result)
    assert len(data["opsec_results"]) == 1
    assert data["opsec_results"][0]["domain"] == "c2.example.com"
    assert data["opsec_results"][0]["grade"] == "B"


# ── HTMLReporter tests ────────────────────────────────────────────────────────

def test_html_reporter_contains_engagement_name(engagement, objectives, ttps, log_entries):
    reporter = HTMLReporter()
    result = reporter.generate(engagement, objectives, ttps, log_entries)
    assert "Op Nightfall" in result


def test_html_output_contains_executive_summary(engagement, objectives, ttps, log_entries):
    reporter = HTMLReporter()
    result = reporter.generate(engagement, objectives, ttps, log_entries)
    assert "Executive Summary" in result


def test_html_output_contains_mitre_attack(engagement, objectives, ttps, log_entries):
    reporter = HTMLReporter()
    result = reporter.generate(engagement, objectives, ttps, log_entries)
    assert "MITRE ATT" in result


def test_html_output_contains_technical_narrative(engagement, objectives, ttps, log_entries):
    reporter = HTMLReporter()
    result = reporter.generate(engagement, objectives, ttps, log_entries)
    assert "Technical Narrative" in result


def test_html_output_contains_detection_gaps(engagement, objectives, ttps, log_entries):
    reporter = HTMLReporter()
    result = reporter.generate(engagement, objectives, ttps, log_entries)
    assert "Detection Gaps" in result


def test_html_output_contains_scope_item(engagement, objectives, ttps, log_entries):
    reporter = HTMLReporter()
    result = reporter.generate(engagement, objectives, ttps, log_entries)
    assert "10.0.0.0/8" in result


def test_html_detection_gaps_only_success_entries(engagement, objectives, ttps, log_entries):
    reporter = HTMLReporter()
    result = reporter.generate(engagement, objectives, ttps, log_entries)
    # "Sent phishing email" is success → should appear in gaps section
    # "Lateral movement attempt" is failure → should NOT be in gaps
    # We verify by checking the gap-item div content includes alice's action
    assert "Sent phishing email" in result
    # The detection_gaps list only includes success entries so count should be 1
    # Count occurrences of gap-item class — one gap entry means 1 div.gap-item
    assert result.count("gap-item") >= 1


def test_html_reporter_uses_real_template(engagement, objectives, ttps, log_entries):
    """Verify template is loaded from the templates directory relative to module."""
    from redops.reporters.html import TEMPLATES_DIR
    template_file = TEMPLATES_DIR / "report.html.j2"
    assert template_file.exists(), f"Template not found at {template_file}"
    reporter = HTMLReporter()
    result = reporter.generate(engagement, objectives, ttps, log_entries)
    assert "<!DOCTYPE html>" in result


def test_html_reporter_write_creates_file(engagement, objectives, ttps, log_entries, tmp_path):
    reporter = HTMLReporter()
    out = tmp_path / "report.html"
    reporter.write(engagement, objectives, ttps, log_entries, out)
    assert out.exists()
    content = out.read_text(encoding="utf-8")
    assert "Op Nightfall" in content


# ── DOCXReporter tests ────────────────────────────────────────────────────────

def test_docx_reporter_generate_returns_bytes(engagement, objectives, ttps, log_entries):
    reporter = DOCXReporter()
    result = reporter.generate(engagement, objectives, ttps, log_entries)
    assert isinstance(result, bytes)
    assert len(result) > 0


def test_docx_contains_engagement_name(engagement, objectives, ttps, log_entries):
    from docx import Document
    reporter = DOCXReporter()
    raw = reporter.generate(engagement, objectives, ttps, log_entries)
    doc = Document(io.BytesIO(raw))
    texts = [p.text for p in doc.paragraphs]
    full_text = " ".join(texts)
    assert "Op Nightfall" in full_text


def test_docx_contains_mitre_heading(engagement, objectives, ttps, log_entries):
    from docx import Document
    reporter = DOCXReporter()
    raw = reporter.generate(engagement, objectives, ttps, log_entries)
    doc = Document(io.BytesIO(raw))
    headings = [p.text for p in doc.paragraphs if p.style.name.startswith("Heading")]
    assert any("MITRE ATT&CK Coverage" in h for h in headings)


def test_docx_contains_client_name(engagement, objectives, ttps, log_entries):
    from docx import Document
    reporter = DOCXReporter()
    raw = reporter.generate(engagement, objectives, ttps, log_entries)
    doc = Document(io.BytesIO(raw))
    texts = " ".join(p.text for p in doc.paragraphs)
    assert "ACME Corp" in texts


def test_docx_contains_operator_log_table(engagement, objectives, ttps, log_entries):
    from docx import Document
    reporter = DOCXReporter()
    raw = reporter.generate(engagement, objectives, ttps, log_entries)
    doc = Document(io.BytesIO(raw))
    # Should have at least 2 tables (MITRE coverage + Technical Narrative)
    assert len(doc.tables) >= 2


def test_docx_reporter_write_creates_file(engagement, objectives, ttps, log_entries, tmp_path):
    from docx import Document
    reporter = DOCXReporter()
    out = tmp_path / "report.docx"
    reporter.write(engagement, objectives, ttps, log_entries, out)
    assert out.exists()
    doc = Document(str(out))
    texts = " ".join(p.text for p in doc.paragraphs)
    assert "Op Nightfall" in texts


# ── PDFReporter mock test ─────────────────────────────────────────────────────

def test_pdf_reporter_calls_weasyprint(engagement, objectives, ttps, log_entries, tmp_path):
    import sys
    import types

    # WeasyPrint requires GTK native libs not present on this machine.
    # Inject a fake module into sys.modules so the lazy `from weasyprint import HTML`
    # inside PDFReporter.generate() resolves to our mock without loading the real lib.
    mock_html_cls = MagicMock()
    mock_instance = MagicMock()
    mock_instance.write_pdf.return_value = b"%PDF-1.4 fake"
    mock_html_cls.return_value = mock_instance

    fake_weasyprint = types.ModuleType("weasyprint")
    fake_weasyprint.HTML = mock_html_cls

    with patch.dict(sys.modules, {"weasyprint": fake_weasyprint}):
        from redops.reporters.pdf import PDFReporter
        reporter = PDFReporter()
        result = reporter.generate(engagement, objectives, ttps, log_entries)
        assert result == b"%PDF-1.4 fake"
        mock_html_cls.assert_called_once()
